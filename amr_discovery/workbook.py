"""Import the supplied long-form AST workbook without inventing provenance."""
from __future__ import annotations

from collections import Counter
import hashlib
import json
from pathlib import Path
import re

import pandas as pd

from .data import IntegrityError


REQUIRED_HEADERS = {
    "#BioSample", "Scientific name", "Antibiotic", "Resistance phenotype",
    "Measurement sign", "MIC (mg/L)", "Laboratory typing platform", "Testing standard",
}
DRUG_ALIASES = {
    "amoxicillin-clavulanic acid": "amoxicillin-clavulanate",
    "amoxicillin clavulanic acid": "amoxicillin-clavulanate",
    "amoxicillin-clavulanate": "amoxicillin-clavulanate",
    "piperacillin-tazobactam": "piperacillin-tazobactam",
}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _clean(value) -> str:
    return "" if value is None else str(value).strip()


def _number_text(value) -> str:
    if value is None or value == "":
        return ""
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise IntegrityError(f"Invalid MIC value in workbook: {value!r}.") from error
    if number <= 0:
        raise IntegrityError(f"MIC value must be positive; found {value!r}.")
    return str(int(number)) if number.is_integer() else format(number, ".12g")


def normalize_rows(headers, rows, sheet_name="Sheet1"):
    """Convert workbook rows to canonical observation records, returning row issues."""
    missing = REQUIRED_HEADERS - set(headers)
    if missing:
        raise IntegrityError(f"Workbook is missing required columns: {sorted(missing)}.")
    positions = {name: headers.index(name) for name in headers}
    output, issues = [], []
    drug_counts, phenotype_counts = Counter(), Counter()
    for source_row, values in enumerate(rows, start=2):
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        source = {name: values[index] if index < len(values) else None for name, index in positions.items()}
        accession = _clean(source.get("#BioSample"))
        species = _clean(source.get("Scientific name"))
        original_drug = _clean(source.get("Antibiotic"))
        original_sir = _clean(source.get("Resistance phenotype")).lower()
        if not accession or not species or not original_drug:
            issues.append({"source_row": source_row, "reason": "missing_isolate_species_or_drug"})
            continue
        if original_sir not in {"resistant", "susceptible", "intermediate", "susceptible, increased exposure", "sdd"}:
            issues.append({"source_row": source_row, "reason": "unsupported_reported_phenotype", "value": original_sir})
            continue
        drug_key = re.sub(r"\s+", " ", original_drug.casefold())
        drug = DRUG_ALIASES.get(drug_key, drug_key.replace("_", "-"))
        mic = _number_text(source.get("MIC (mg/L)"))
        operator = _clean(source.get("Measurement sign"))
        if operator not in {"=", "==", "<", "<=", ">", ">=", "≤", "≥"}:
            issues.append({"source_row": source_row, "reason": "unsupported_measurement_operator", "value": operator})
            continue
        location = _clean(source.get("Location"))
        country = location.split(":", 1)[0].strip() if ":" in location else ""
        row = {
            "isolate_id": accession,
            "source_id": "ncbi_source_unknown",
            "species": species,
            "drug": drug,
            "measurement": mic,
            "operator": operator,
            "units": "mg/L",
            "method": "reported AST",
            "standard": _clean(source.get("Testing standard")),
            "standard_version": "",
            "reported_sir": original_sir,
            "evidence_class": "measured_authorized",
            "evidence_reference": "User-supplied workbook; source study not identified",
            "accession": accession,
            "source_row": f"{sheet_name}!{source_row}",
            "country": country,
            "collection_date": "",
            "host": "",
            "specimen": _clean(source.get("Isolation source")),
            "lab_id": "",
            "platform": _clean(source.get("Laboratory typing platform")),
            "vendor": _clean(source.get("Vendor")),
            "record_create_date": _clean(source.get("Create date")),
            "original_drug": original_drug,
            "original_location": location,
            "method_version_or_reagent": _clean(source.get("Laboratory typing method version or reagent")),
        }
        output.append(row)
        drug_counts[drug] += 1
        phenotype_counts[f"{drug}|{original_sir}"] += 1
    return output, issues, {"drug_rows": dict(sorted(drug_counts.items())),
                            "drug_phenotype_rows": dict(sorted(phenotype_counts.items()))}


def import_asts_workbook(source, output_dir):
    """Write normalized observations and a manifest under the ignored private-data folder."""
    source = Path(source).expanduser().resolve()
    destination = Path(output_dir).expanduser().resolve()
    private_root = Path("data/private").resolve()
    try:
        destination.relative_to(private_root)
    except ValueError as error:
        raise IntegrityError("Workbook output must be inside ignored data/private/.") from error
    if destination.exists():
        raise FileExistsError(f"Output exists; choose a new private output directory: {destination}")
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise IntegrityError("Install workbook support with `python -m pip install -e '.[workbooks]'`.") from error
    if source.suffix.lower() != ".xlsx":
        raise IntegrityError("The AST workbook importer accepts .xlsx files.")
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [_clean(value) for value in next(iterator)]
        except StopIteration as error:
            raise IntegrityError("Workbook sheet is empty.") from error
        rows, issues, counts = normalize_rows(headers, iterator, sheet.title)
    finally:
        workbook.close()
    destination.mkdir(parents=True, exist_ok=False)
    observations = destination / "observations.csv"
    pd.DataFrame(rows).to_csv(observations, index=False, lineterminator="\n")
    manifest = {
        "schema_version": 1,
        "status": "PREPARED_NOT_AUDITED",
        "input_name": source.name,
        "input_sha256": _sha256(source),
        "output_sha256": _sha256(observations),
        "sheet": sheet.title,
        "source_rows_read": sum(counts["drug_rows"].values()) + len(issues),
        "normalized_rows": len(rows),
        "issue_rows": len(issues),
        "issues": issues,
        "counts": counts,
        "unknown_provenance": ["study/source identity", "standard version", "collection date",
                               "patient linkage", "laboratory identity", "host confirmation"],
        "notes": ["Workbook records are retained as user-authorized reported AST.",
                  "Create date is retained separately and is not treated as collection date.",
                  "Location is retained; country is populated only when explicitly prefixed.",
                  "No MIC values or reported phenotypes were altered."],
    }
    (destination / "import_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return manifest
